import os
import io
import csv
import secrets
from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask import (Blueprint, render_template_string, redirect, url_for, request,
                    session, flash, jsonify, Response)

import wiki_auth
from flask_app import (MANUAL_ADMIN_ID, MANUAL_ADMIN_USERNAME, MANUAL_LOGIN_ID, MANUAL_LOGIN_PASS,
                  WIKI_ROOT_ADMIN_USERNAME,
                  ROLE_NONE, ROLE_COLLAB, ROLE_ADMIN, RECAPTCHA_SITE_KEY, BOTS_DIR,
                  get_db, get_text, verify_recaptcha, get_all_bots, get_script_status,
                  login_required, check_role, require_api_key, csrf, status,
                  launch_script_core, stop_current_script)
from templates import (GLASS_CSS, GATE_HTML, LOGIN_MANUAL_HTML, ACCOUNT_HTML,
                        DASHBOARD_HTML, HISTORY_HTML)

auth_bp = Blueprint("auth", __name__)
dashboard_bp = Blueprint("dashboard", __name__)
api_bp = Blueprint("api", __name__)


# ============================================================
# AUTH
# ============================================================

@auth_bp.route("/security-gate")
def security_gate():
    return render_template_string(GATE_HTML, site_key=RECAPTCHA_SITE_KEY, glass_css=GLASS_CSS)


@auth_bp.route("/verify-gate", methods=["POST"])
def verify_gate():
    captcha_response = request.form.get('g-recaptcha-response')
    if verify_recaptcha(captcha_response):
        session['captcha_passed'] = True
        return redirect(url_for('dashboard.index'))
    flash("Veuillez valider le captcha correctement.")
    return redirect(url_for('auth.security_gate'))


@auth_bp.route("/set_lang/<code>")
def set_language(code):
    if code not in ['fr', 'en']:
        code = 'fr'
    session['lang'] = code
    if 'user_id' in session:
        try:
            db = get_db()
            db.execute("UPDATE users SET lang = ? WHERE wiki_id = ?", (code, session['user_id']))
            db.commit()
        except Exception:
            pass
    return redirect(request.referrer or url_for('dashboard.index'))


@auth_bp.route("/admin/login")
def manual_login_page():
    return render_template_string(LOGIN_MANUAL_HTML, glass_css=GLASS_CSS)


@auth_bp.route("/admin/login_post", methods=['POST'])
def manual_login_post():
    username = request.form.get('username')
    password = request.form.get('password')
    if username == MANUAL_LOGIN_ID and password == MANUAL_LOGIN_PASS:
        db = get_db()
        existing_lang = db.execute("SELECT lang FROM users WHERE wiki_id = ?", (MANUAL_ADMIN_ID,)).fetchone()
        current_lang = existing_lang['lang'] if existing_lang else 'fr'
        db.execute("""INSERT OR REPLACE INTO users (wiki_id, username, avatar, role, is_banned, lang)
                        VALUES (?, ?, ?, ?, 0, ?)""",
                   (MANUAL_ADMIN_ID, MANUAL_ADMIN_USERNAME,
                    "https://ui-avatars.com/api/?name=Admin+Bot&background=ff0000&color=fff", ROLE_ADMIN, current_lang))
        db.commit()
        session.permanent = True
        session['user_id'] = MANUAL_ADMIN_ID
        session['username'] = MANUAL_ADMIN_USERNAME
        session['lang'] = current_lang
        flash("Connexion Admin Manuelle Réussie.")
        return redirect(url_for('dashboard.index'))
    flash(get_text('error_manual_login'))
    return redirect(url_for('auth.manual_login_page'))


@auth_bp.route("/login")
def login_wiki():
    # Si une liaison Discord est en cours (voir /discord/link/<token>), on la garde
    # en session le temps du round-trip OAuth pour la reprendre après connexion.
    verifier, challenge = wiki_auth.generate_pkce_pair()
    state = secrets.token_urlsafe(24)
    session['wiki_oauth_state'] = state
    session['wiki_oauth_verifier'] = verifier
    return redirect(wiki_auth.build_authorize_url(state, challenge))


@auth_bp.route("/oauth/wiki/callback")
@csrf.exempt
def callback_wiki():
    code = request.args.get('code')
    state = request.args.get('state')
    expected_state = session.pop('wiki_oauth_state', None)
    verifier = session.pop('wiki_oauth_verifier', None)

    if not code or not state or not verifier or state != expected_state:
        flash("Échec de la connexion Vikidia (état invalide).")
        return redirect(url_for('dashboard.index'))

    token_data = wiki_auth.exchange_code_for_token(code, verifier)
    if not token_data:
        flash("Échec de la connexion Vikidia (échange du code impossible).")
        return redirect(url_for('dashboard.index'))

    profile = wiki_auth.fetch_profile(token_data['access_token'])
    if not profile:
        flash("Échec de la connexion Vikidia (profil introuvable).")
        return redirect(url_for('dashboard.index'))

    wiki_id = str(profile.get('sub') or profile.get('id') or profile['username'])
    username = profile['username']
    avatar_url = f"https://ui-avatars.com/api/?name={username}&background=0093E9&color=fff"

    db = get_db()
    existing_user = db.execute("SELECT * FROM users WHERE wiki_id = ?", (wiki_id,)).fetchone()
    role, lang = ROLE_NONE, 'fr'
    if username == WIKI_ROOT_ADMIN_USERNAME:
        role = ROLE_ADMIN

    if existing_user:
        role, lang = existing_user['role'], existing_user['lang'] or 'fr'
        if existing_user['is_banned']:
            flash(f"Compte banni. Raison : {existing_user['ban_reason'] or 'Non spécifiée'}")
            return redirect(url_for('dashboard.index'))

    db.execute("""INSERT OR REPLACE INTO users (wiki_id, username, avatar, role, is_banned, lang, ban_reason, discord_id)
                  VALUES (?, ?, ?, ?, COALESCE((SELECT is_banned FROM users WHERE wiki_id=?), 0), ?,
                          (SELECT ban_reason FROM users WHERE wiki_id=?),
                          (SELECT discord_id FROM users WHERE wiki_id=?))""",
               (wiki_id, username, avatar_url, role, wiki_id, lang, wiki_id, wiki_id))
    db.commit()

    session['user_id'], session['username'], session['lang'] = wiki_id, username, lang

    # Reprise d'une liaison Discord démarrée avant la connexion (/discord/link/<token>).
    pending_token = session.pop('pending_discord_link', None)
    if pending_token:
        return redirect(url_for('auth.discord_link_confirm', token=pending_token))

    return redirect(url_for('dashboard.index'))


# ============================================================
# LIAISON DE COMPTE DISCORD <-> VIKIDIA
# ============================================================
# Le robot Discord appelle POST /api/discord/link/start pour obtenir une URL à usage
# unique. L'utilisateur l'ouvre, se connecte via Vikidia OAuth si besoin, puis confirme
# la liaison ici. Ensuite, GET /api/discord/permissions?discord_id=... permet au robot
# de vérifier les droits (voir auth_check.py côté bot).

DISCORD_LINK_TOKEN_TTL_MINUTES = 15


@auth_bp.route("/discord/link/<token>")
def discord_link_confirm(token):
    db = get_db()
    link = db.execute("SELECT * FROM discord_links WHERE token = ?", (token,)).fetchone()
    if not link or link['used']:
        flash("Ce lien de liaison Discord est invalide ou a déjà été utilisé.")
        return redirect(url_for('dashboard.index'))

    created_at = datetime.strptime(link['created_at'], "%Y-%m-%d %H:%M:%S")
    if datetime.now() - created_at > timedelta(minutes=DISCORD_LINK_TOKEN_TTL_MINUTES):
        flash("Ce lien de liaison Discord a expiré, relance /auth sur Discord.")
        return redirect(url_for('dashboard.index'))

    if 'user_id' not in session:
        # On garde le token en session pour reprendre la liaison juste après le login.
        session['pending_discord_link'] = token
        return redirect(url_for('auth.login_wiki'))

    db.execute("UPDATE users SET discord_id = ? WHERE wiki_id = ?", (link['discord_id'], session['user_id']))
    db.execute("UPDATE discord_links SET used = 1 WHERE token = ?", (token,))
    db.commit()
    flash(f"Compte Discord ({link['discord_username']}) lié avec succès à votre compte Vikidia.")
    return redirect(url_for('auth.account'))


@auth_bp.route("/logout")
def logout():
    session.clear()
    return redirect(url_for('dashboard.index'))


@auth_bp.route("/account")
@login_required
def account():
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE wiki_id = ?", (session['user_id'],)).fetchone()
    return render_template_string(ACCOUNT_HTML, user=user, glass_css=GLASS_CSS)


# ============================================================
# DASHBOARD
# ============================================================

@dashboard_bp.route("/")
def index():
    db = get_db()
    row = db.execute("SELECT value FROM settings WHERE key='lock_launch'").fetchone()
    locked = row['value'] if row else '0'
    user_role = ROLE_NONE
    if 'user_id' in session:
        u = db.execute("SELECT role FROM users WHERE wiki_id=?", (session['user_id'],)).fetchone()
        if u:
            user_role = u['role']

    all_files = get_all_bots()
    available_scripts = []
    for f in all_files:
        is_act = get_script_status(f)
        if user_role == ROLE_ADMIN or is_act == 1:
            available_scripts.append({"filename": f, "is_active": is_act})

    return render_template_string(
        DASHBOARD_HTML, running=status["running"], script_name=status.get("script_name") or "Inactif",
        logs=status["live_output"][-50:], locked=locked, role=user_role, glass_css=GLASS_CSS,
        available_scripts=available_scripts
    )


@dashboard_bp.route("/start", methods=["POST"])
@login_required
def start_script():
    db = get_db()
    locked_val = db.execute("SELECT value FROM settings WHERE key='lock_launch'").fetchone()['value']
    user = db.execute("SELECT * FROM users WHERE wiki_id=?", (session['user_id'],)).fetchone()

    if not user or user['is_banned']:
        flash("Compte banni.")
        return redirect(url_for('dashboard.index'))

    if user['role'] not in (ROLE_COLLAB, ROLE_ADMIN):
        # Rôle "None" : vérification en direct du statut Autopatrolleur sur au moins
        # une des versions linguistiques de Vikidia prises en charge. Aucune promotion
        # de rôle n'est enregistrée en base : le contrôle est refait à chaque lancement.
        if not wiki_auth.is_autopatrolled_anywhere(user['username']):
            flash(get_text('error_not_autopatrolled'))
            return redirect(url_for('dashboard.index'))

    if locked_val == '1' and user['role'] != ROLE_ADMIN:
        flash("Verrouillé par l'Admin.")
        return redirect(url_for('dashboard.index'))

    if not status["running"]:
        choice = request.form.get("choice")
        p = os.path.join(BOTS_DIR, choice)

        if user['role'] != ROLE_ADMIN and get_script_status(choice) == 0:
            flash("Ce script a été désactivé par l'administrateur.")
            return redirect(url_for('dashboard.index'))

        if os.path.exists(p):
            cmd_args = []
            if "portal.py" in choice or choice == "Portail":
                arg_lang = request.form.get("arg_lang")
                arg_cat = request.form.get("arg_cat")
                arg_portal = request.form.get("arg_portal")
                if not (arg_lang and arg_cat and arg_portal):
                    return redirect(url_for('dashboard.index'))
                cmd_args = ["--lang", arg_lang, "--cat", arg_cat, "--portal", arg_portal]

            launch_script_core(choice, p, args=cmd_args, user_name=session.get('username'))
    return redirect(url_for("dashboard.index"))


@dashboard_bp.route("/stop", methods=["POST"])
@check_role([ROLE_COLLAB, ROLE_ADMIN])
def stop_script():
    stop_current_script(user_name=session.get('username'))
    return redirect(url_for("dashboard.index"))


# ============================================================
# HISTORIQUE DES LOGS + EXPORTS
# ============================================================

def _get_logs_for_export(f_search=None):
    query = "SELECT id, date, script, message FROM logs WHERE 1=1"
    params = []
    if f_search:
        query += " AND message LIKE ?"
        params.append("%" + f_search + "%")
    query += " ORDER BY id DESC"
    db = get_db()
    return db.execute(query, params).fetchall()


@dashboard_bp.route("/history")
def history():
    f_search = request.args.get('search')
    query = "SELECT * FROM logs WHERE 1=1"
    params = []
    if f_search:
        query += " AND message LIKE ?"
        params.append(f"%{f_search}%")
    query += " ORDER BY id DESC LIMIT 500"
    db = get_db()
    rows = db.execute(query, params).fetchall()

    count_query = "SELECT COUNT(*) FROM logs WHERE 1=1"
    if f_search:
        count_query += " AND message LIKE ?"
    total_count = db.execute(count_query, params).fetchone()[0]

    return render_template_string(HISTORY_HTML, rows=rows, request=request, glass_css=GLASS_CSS, total_count=total_count)


@dashboard_bp.route("/history/export/excel")
def export_logs_excel():
    f_search = request.args.get('search')
    rows = _get_logs_for_export(f_search)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logs BotJanus"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0B132B", end_color="0B132B", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = ["ID", "Date", "Script", "Message"]
    col_widths = [8, 22, 25, 80]
    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22

    fill_light = PatternFill(start_color="EAF4FB", end_color="EAF4FB", fill_type="solid")
    fill_dark = PatternFill(start_color="D0E8F2", end_color="D0E8F2", fill_type="solid")

    for row_idx, row in enumerate(rows, start=2):
        fill = fill_light if row_idx % 2 == 0 else fill_dark
        values = [row[0], row[1], row[2], row[3]]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx == 4))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:D1"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = "logs_botjanus_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx"
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=" + filename}
    )


@dashboard_bp.route("/history/export/csv")
def export_logs_csv():
    f_search = request.args.get('search')
    rows = _get_logs_for_export(f_search)

    si = io.StringIO()
    cw = csv.writer(si, delimiter=';')
    cw.writerow(['ID', 'Date', 'Script', 'Message'])
    for row in rows:
        cw.writerow([row[0], row[1], row[2], row[3]])

    filename = "logs_botjanus_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".csv"
    return Response(
        si.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=" + filename}
    )


# ============================================================
# API DISCORD
# ============================================================

# ============================================================
# API DISCORD
# ============================================================

@api_bp.route('/api/status', methods=['GET'])
@csrf.exempt
@require_api_key
def api_status():
    return jsonify({
        "running": status["running"],
        "script_name": status.get("script_name") or "Inactif",
        "last_activity": status["last_activity"].isoformat() if status["last_activity"] else None
    })


@api_bp.route('/api/status_json')
def status_json():
    db = get_db()
    row = db.execute("SELECT value FROM settings WHERE key='lock_launch'").fetchone()
    locked = row['value'] if row else '0'
    user_role = ROLE_NONE
    if 'user_id' in session:
        u = db.execute("SELECT role FROM users WHERE wiki_id=?", (session['user_id'],)).fetchone()
        if u:
            user_role = u['role']
    return jsonify({
        "running": status["running"],
        "script_name": status.get("script_name") or "Inactif",
        "locked": locked,
        "role": user_role
    })


@api_bp.route('/api/start', methods=['POST'])
@csrf.exempt
@require_api_key
def api_start():
    data = request.get_json()
    if not data:
        return jsonify({"error": "Aucune donnée reçue"}), 400
    choice = data.get('choice')
    discord_username = data.get('username', 'Discord')
    if not choice:
        return jsonify({"error": "Paramètre 'choice' manquant"}), 400
    if status["running"]:
        return jsonify({"error": "Un script est déjà en cours"}), 409

    p = os.path.join(BOTS_DIR, choice)
    if os.path.exists(p):
        cmd_args = []
        if choice == "portal.py" or choice == "Portail":
            cmd_args = ["--lang", data.get("arg_lang", "fr"), "--cat", data.get("arg_cat", ""), "--portal", data.get("arg_portal", "")]
        success = launch_script_core(choice, p, args=cmd_args, user_name=f"Discord ({discord_username})")
        if success:
            return jsonify({"success": True, "message": f"Script {choice} lancé."}), 200
    return jsonify({"error": "Fichier introuvable"}), 404


@api_bp.route('/api/stop', methods=['POST'])
@csrf.exempt
@require_api_key
def api_stop():
    if not status["running"] or not status["process"]:
        return jsonify({"error": "Aucun script en cours"}), 400
    try:
        stop_current_script(user_name="Discord")
        return jsonify({"success": True, "message": "Script arrêté."}), 200
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@api_bp.route('/api/logs', methods=['GET'])
@csrf.exempt
@require_api_key
def api_logs():
    limit = request.args.get('limit', 50, type=int)
    return jsonify({"logs": status["live_output"][-limit:]})


@api_bp.route('/api/scripts', methods=['GET'])
@csrf.exempt
@require_api_key
def api_scripts():
    scripts = []
    for filename in get_all_bots():
        scripts.append({"filename": filename, "is_active": bool(get_script_status(filename))})
    return jsonify({"scripts": scripts})


@api_bp.route("/api/live_logs")
def live_logs():
    return "\n".join([f"<div>{line}</div>" for line in status["live_output"][-100:]])


# ============================================================
# API DISCORD — LIAISON DE COMPTE & PERMISSIONS (utilisées par auth_check.py côté bot)
# ============================================================

@api_bp.route('/api/discord/link/start', methods=['POST'])
@csrf.exempt
@require_api_key
def api_discord_link_start():
    data = request.get_json(silent=True) or {}
    discord_id = data.get('discord_id')
    discord_username = data.get('discord_username', 'Discord')
    if not discord_id:
        return jsonify({"error": "Paramètre 'discord_id' manquant"}), 400

    token = secrets.token_urlsafe(24)
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    db = get_db()
    db.execute(
        "INSERT INTO discord_links (token, discord_id, discord_username, created_at, used) VALUES (?, ?, ?, ?, 0)",
        (token, str(discord_id), discord_username, created_at)
    )
    db.commit()
    return jsonify({"url": url_for('auth.discord_link_confirm', token=token, _external=True)}), 200


@api_bp.route('/api/discord/permissions', methods=['GET'])
@csrf.exempt
@require_api_key
def api_discord_permissions():
    discord_id = request.args.get('discord_id')
    if not discord_id:
        return jsonify({"error": "Paramètre 'discord_id' manquant"}), 400

    db = get_db()
    user = db.execute("SELECT * FROM users WHERE discord_id = ?", (str(discord_id),)).fetchone()
    if not user:
        return jsonify({"linked": False, "authorized": False, "reason": "not_linked"})

    if user['is_banned']:
        return jsonify({"linked": True, "authorized": False, "reason": "banned",
                         "username": user['username']})

    if user['role'] in (ROLE_COLLAB, ROLE_ADMIN):
        return jsonify({"linked": True, "authorized": True, "role": user['role'],
                         "username": user['username']})

    # Rôle "None" lié : même règle que sur le dashboard web, vérification en direct
    # (aucune promotion de rôle enregistrée).
    if wiki_auth.is_autopatrolled_anywhere(user['username']):
        return jsonify({"linked": True, "authorized": True, "role": ROLE_NONE,
                         "via": "autopatrol", "username": user['username']})

    return jsonify({"linked": True, "authorized": False, "reason": "not_autopatrolled",
                     "username": user['username']})
